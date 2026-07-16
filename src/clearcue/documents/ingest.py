from __future__ import annotations

import re
from pathlib import Path


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def _normalise(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    paragraphs: list[str] = []
    current: list[str] = []
    for line in lines:
        if line:
            current.append(line)
        elif current:
            paragraphs.append(" ".join(current))
            current = []
    if current:
        paragraphs.append(" ".join(current))
    return "\n\n".join(paragraphs).strip()


def extract_text(path: str | Path) -> str:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(source))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document

        document = Document(str(source))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        text = source.read_text(encoding="utf-8", errors="replace")

    cleaned = _normalise(text)
    if not cleaned:
        raise ValueError("No readable text was found in this file.")
    return cleaned


def chunk_text(text: str, words_per_chunk: int = 190, overlap: int = 35) -> list[str]:
    """Split context into compact overlapping chunks without external NLP packages."""
    words = text.split()
    if not words:
        return []
    words_per_chunk = max(60, words_per_chunk)
    overlap = min(max(0, overlap), words_per_chunk // 2)
    step = words_per_chunk - overlap
    chunks: list[str] = []
    for start in range(0, len(words), step):
        piece = words[start : start + words_per_chunk]
        if not piece:
            break
        chunks.append(" ".join(piece))
        if start + words_per_chunk >= len(words):
            break
    return chunks

